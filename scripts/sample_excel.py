from pathlib import Path
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
import sys
sys.path.insert(0, str(Path(__file__).resolve().parents[1] / "backend"))
from app.schemas import COLUMNS

root = Path(__file__).resolve().parents[1]
output = root / "data" / "示例单元测试报告.xlsx"
output.parent.mkdir(exist_ok=True)
book = Workbook()
for n, domain in enumerate(["射频", "协议"]):
    sheet = book.active if n == 0 else book.create_sheet()
    sheet.title = domain
    sheet.append(COLUMNS)
    sheet.append(["TC-001", "基础功能验证", "测试小项 A", "输出值", "常温", "参考门限", "dB", "10", "12", "仅用于演示", "https://intranet.example/evidence/1.jpg"])
    sheet.append(["", "", "测试小项 B", "输出值", "常温", "参考门限", "dB", "10", "8", "", "https://intranet.example/evidence/2.jpg"])
    sheet.append(["TC-002", "无附件文字项", "", "", "", "", "", "", "正常", "", ""])
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet["K2"].hyperlink = sheet["K2"].value
    sheet["K2"].value = "查看附件"
    sheet.freeze_panes = "A2"
    for cell in sheet[1]:
        cell.font = Font(color="FFFFFF", bold=True)
        cell.fill = PatternFill("solid", fgColor="146D69")
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 20
book.save(output)
print(output)
