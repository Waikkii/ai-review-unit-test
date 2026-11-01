from io import BytesIO
from zipfile import ZipFile, BadZipFile
import re
from openpyxl import load_workbook
from .schemas import COLUMNS
from .config import settings

def attachment_urls(value):
    # 文本支持多个 URL（逗号、空格、换行分隔），不执行 Excel 公式。
    return re.findall(r"https?://[^\s<>\"，,;；]+", value, flags=re.I)

def detail_header(sheet):
    """识别测试详情双层表头，兼容单行模板，排除上方测试总结。"""
    sections = [cell.row for row in sheet.iter_rows() for cell in row
                if cell.value == "测试详情"]
    if len(sections) > 1:
        raise ValueError(f"{sheet.title}: 存在多个测试详情区域，暂不支持")
    start = sections[0] + 1 if sections else 1
    end = min(sheet.max_row, start + 29)
    for bottom in range(start, end + 1):
        for top in range(bottom, max(start, bottom - 1) - 1, -1):
            positions = {name: set() for name in ["测试对象", *COLUMNS]}
            for row in sheet.iter_rows(min_row=top, max_row=bottom):
                for cell in row:
                    name = str(cell.value).strip() if cell.value is not None else ""
                    if name in positions:
                        positions[name].add(cell.column)
            if all(len(positions[name]) == 1 for name in COLUMNS) and len(positions["测试对象"]) <= 1:
                return bottom, {name: next(iter(cols)) for name, cols in positions.items() if cols}
    raise ValueError(f"{sheet.title}: 未找到测试详情完整、唯一的 11 个字段（支持两行表头）")


def parse_excel(content):
    try:
        with ZipFile(BytesIO(content)) as archive:
            infos = archive.infolist()
            if len(infos) > 10000 or sum(i.file_size for i in infos) > 150 * 1024 * 1024:
                raise ValueError("Excel 解压后过大")
        workbook = load_workbook(BytesIO(content), data_only=False, keep_links=False)
    except (BadZipFile, OSError, KeyError) as exc:
        raise ValueError("无法读取有效的 .xlsx 文件") from exc
    parsed = []
    try:
        for sheet in workbook:
            if sheet.max_row * sheet.max_column > 2_000_000:
                raise ValueError(f"{sheet.title} 工作表范围过大")
            if not any(c.value is not None for row in sheet.iter_rows(max_row=min(sheet.max_row, 30)) for c in row):
                continue
            if sheet.title.strip() in {"封面", "目录"}:
                continue
            hrow, columns = detail_header(sheet)
            merged = {}
            for area in sheet.merged_cells.ranges:
                # 横向合并是说明/分组行，不将“/”扩散为实测值。
                if area.min_row <= hrow or area.min_col != area.max_col:
                    continue
                top = sheet.cell(area.min_row, area.min_col)
                for row in range(area.min_row, area.max_row + 1):
                    for col in range(area.min_col, area.max_col + 1):
                        merged[row, col] = top
            for row in range(hrow + 1, sheet.max_row + 1):
                value_cell = merged.get((row, columns["值"]), sheet.cell(row, columns["值"]))
                if value_cell.value is None or not str(value_cell.value).strip():
                    continue
                values = {}
                for name, col in columns.items():
                    cell = merged.get((row, col), sheet.cell(row, col))
                    value = cell.value
                    if name == "附件链接" and cell.hyperlink and cell.hyperlink.target:
                        value = cell.hyperlink.target
                    elif cell.data_type == "f":
                        # HYPERLINK 常量可读取；其他公式拒绝猜测缓存值。
                        match = re.fullmatch(r'(?i)=HYPERLINK\(\s*"([^"]+)"\s*[,;]\s*"[^"]*"\s*\)', str(value))
                        if name == "附件链接" and match:
                            value = match.group(1)
                        else:
                            raise ValueError(f"{sheet.title} 第 {row} 行 {name}: 请将公式转成值后上传")
                    values[name] = "" if value is None else str(value)
                parsed.append({"domain": sheet.title, "row_number": row, "fields": values})
                if len(parsed) > settings.max_items:
                    raise ValueError("待审查项超过配置上限")
    finally:
        workbook.close()
    if not parsed:
        raise ValueError("未发现测试详情中“值”非空的测试项")
    return parsed

def safe_cell(value):
    if isinstance(value, str) and value.startswith(("=", "+", "-", "@")):
        return "'" + value
    return value

def export_excel(original, run, results):
    workbook = load_workbook(BytesIO(original), keep_links=False)
    sheet = workbook.create_sheet("AI审查_" + run.id[:8])
    sheet.append(["领域", "原始行号", "测试对象", *COLUMNS, "审查状态", "问题说明", "判断依据", "修改建议",
                  "审查批次", "报告版本", "模型", "规则版本", "运行模式"])
    for result in results:
        values = [result.domain, result.row_number, result.fields.get("测试对象", ""), *[result.fields.get(c, "") for c in COLUMNS],
                  result.status, result.explanation, "\n".join(result.evidence), result.suggestion,
                  run.id, run.revision, run.model, run.rule_version, run.mode]
        sheet.append([safe_cell(v) for v in values])
    sheet.freeze_panes = "A2"
    sheet.auto_filter.ref = sheet.dimensions
    for col in sheet.columns:
        sheet.column_dimensions[col[0].column_letter].width = 22
    out = BytesIO()
    workbook.save(out)
    workbook.close()
    return out.getvalue()
