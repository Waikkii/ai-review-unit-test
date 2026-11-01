from io import BytesIO
from datetime import timedelta
from concurrent.futures import ThreadPoolExecutor
from openpyxl import Workbook, load_workbook
from sqlalchemy import select
from app.schemas import COLUMNS, ReviewVerdict
from app.db import Session, now
from app.models import User, Run, Result, Audit
from app.excel import parse_excel
from app.worker import process_once

def workbook_bytes():
    book = Workbook()
    sheet = book.active
    sheet.title = "射频"
    sheet.append(COLUMNS)
    sheet.append(["TC1", "功能测试", "小项 A", "功率", "常温", "大于", "dB", "10", "12", "", "查看图片"])
    sheet["K2"].hyperlink = "https://intranet.example/a.jpg"
    sheet.append(["", "", "小项 B", "功率", "常温", "大于", "dB", "10", "8", "", "https://intranet.example/b.jpg"])
    sheet.merge_cells("A2:A3")
    sheet.merge_cells("B2:B3")
    sheet.append(["TC2", "无附件", "", "", "", "", "", "", "", "", ""])
    second = book.create_sheet("协议")
    second.append(COLUMNS)
    second.append(["TC3", "协议检查", "", "", "", "", "", "", "正常", "", "https://intranet.example/c.jpg"])
    out = BytesIO()
    book.save(out)
    book.close()
    return out.getvalue()

def upload(client, h):
    response = client.post("/api/reports", headers=h, files={"file": ("sample.xlsx", workbook_bytes())})
    assert response.status_code == 201, response.text
    return response.json()

def change(domain="射频", action="create", key="power", version=0, value="功率"):
    return {"action":action, "domain":domain, "key":key, "payload":{"words":[value]},
            "base_version":version, "reason":"完善识别关键词"}

def test_parse_merged_hyperlinks_and_multi_sheet():
    rows = parse_excel(workbook_bytes())
    assert len(rows) == 3
    assert rows[1]["fields"]["用例编号"] == "TC1"
    assert rows[0]["fields"]["附件链接"] == "https://intranet.example/a.jpg"
    assert rows[2]["domain"] == "协议"

def test_auth_and_role_injection(client, headers):
    assert client.get("/api/reports").status_code == 401
    assert client.post("/api/auth/login",json={"username":"admin","password":"wrong"}).status_code == 401
    assert client.post("/api/auth/login",json={"username":"demo","password":"test-password-123","role":"admin"}).status_code == 422
    assert client.get("/api/users",headers=headers()).status_code == 403

def test_upload_rejects_invalid_file(client, headers):
    h = headers()
    assert client.post("/api/reports",headers=h,files={"file":("bad.xlsx",b"invalid")}).status_code == 400
    assert client.post("/api/reports",headers=h,files={"file":("bad.xls",b"invalid")}).status_code == 400

def test_owner_visibility_and_admin_readonly(client, headers):
    admin, user = headers("admin"), headers()
    report = upload(client, admin)
    rid = report["id"]
    assert client.get("/api/reports",headers=user).json() == []
    for suffix in ["items","original","runs"]:
        assert client.get(f"/api/reports/{rid}/{suffix}",headers=user).status_code == 404
    assert client.post(f"/api/reports/{rid}/runs",headers=user).status_code == 404
    own = upload(client,user)
    item = client.get(f'/api/reports/{own["id"]}/items',headers=admin).json()[0]
    assert client.patch('/api/items/'+item["id"],headers=admin,
        json={"version":1,"fields":item["current"],"reason":"不允许"}).status_code == 404

def test_run_progress_export_and_immutable_snapshots(client, headers):
    h = headers()
    report = upload(client,h)
    rid = report["id"]
    original = client.get(f"/api/reports/{rid}/original",headers=h).content
    run = client.post(f"/api/reports/{rid}/runs",headers=h).json()
    assert client.post(f"/api/reports/{rid}/runs",headers=h).status_code == 409
    item = client.get(f"/api/reports/{rid}/items",headers=h).json()[0]
    fields = {**item["current"], "备注":"=1+1"}
    response = client.patch('/api/items/'+item["id"],headers=h,json={"version":1,"fields":fields,"reason":"补充备注"})
    assert response.status_code == 200
    assert client.patch('/api/items/'+item["id"],headers=h,
        json={"version":1,"fields":fields,"reason":"旧版本"}).status_code == 409
    assert process_once()
    detail = client.get('/api/runs/'+run["id"],headers=h).json()
    assert (detail["state"],detail["completed"],detail["failed"]) == ("completed",3,0)
    results = client.get('/api/runs/'+run["id"]+'/results',headers=h).json()
    assert all(r["status"] == "needs_review" and r["raw"]["demo"] for r in results)
    assert next(r for r in results if r["item_id"]==item["id"])["fields"]["备注"] == ""
    assert client.get(f"/api/reports/{rid}/original",headers=h).content == original
    exported = client.get('/api/runs/'+run["id"]+'/export',headers=h)
    book = load_workbook(BytesIO(exported.content))
    assert book.sheetnames[:2] == ["射频","协议"]
    assert len(book.worksheets[-1]["A"]) == 4
    book.close()
    second = client.post(f"/api/reports/{rid}/runs",headers=h).json()
    assert second["revision"] == 2
    process_once()
    book = load_workbook(BytesIO(client.get('/api/runs/'+second["id"]+'/export',headers=h).content))
    assert any(cell.value == "'=1+1" for row in book.worksheets[-1] for cell in row)
    book.close()

def test_keyword_approval_conflict_and_audit(client, headers):
    user, admin = headers(), headers("admin")
    create = client.post('/api/keyword-requests',headers=user,json=change()).json()
    assert create["state"] == "pending"
    assert client.get('/api/keywords',headers=user).json() == []
    assert client.post('/api/keyword-requests/'+create["id"]+'/decision',headers=user,
                       json={"approve":True,"note":"越权"}).status_code == 403
    assert client.post('/api/keyword-requests/'+create["id"]+'/decision',headers=admin,
                       json={"approve":True,"note":"同意新增"}).status_code == 200
    assert client.get('/api/keywords',headers=user).json()[0]["version"] == 1
    pending = [client.post('/api/keyword-requests',headers=user,
        json=change(action="update",version=1,value=v)).json()["id"] for v in ["值 A","值 B"]]
    assert client.post('/api/keyword-requests/'+pending[0]+'/decision',headers=admin,
                       json={"approve":True,"note":"批准"}).status_code == 200
    assert client.post('/api/keyword-requests/'+pending[1]+'/decision',headers=admin,
                       json={"approve":True,"note":"过期版本"}).status_code == 409
    assert client.get('/api/keywords',headers=user).json()[0]["payload"]["words"] == ["值 A"]
    direct = client.post('/api/keyword-requests',headers=admin,json=change(action="delete",version=2))
    assert direct.status_code == 201 and direct.json()["state"] == "approved"
    assert client.get('/api/keywords',headers=user).json() == []
    assert any(a["action"] == "keyword.delete" and a["before"]["version"] == 2
               for a in client.get('/api/audit',headers=admin).json())

def test_permissions_apply_to_existing_tokens(client, headers):
    admin,user = headers("admin"),headers()
    target = next(u for u in client.get('/api/users',headers=admin).json() if u["username"]=="demo")
    assert client.patch('/api/users/'+target["id"],headers=admin,
        json={"role":"user","active":False,"reason":"停用测试"}).status_code == 200
    assert client.get('/api/reports',headers=user).status_code == 401
    assert client.post('/api/auth/login',json={"username":"demo","password":"test-password-123"}).status_code == 403

def test_worker_failure_recovery_and_lease(client, headers, monkeypatch):
    import app.worker as worker
    h=headers()
    report=upload(client,h)
    run=client.post('/api/reports/'+report["id"]+'/runs',headers=h).json()
    with Session() as db:
        row=db.get(Run,run["id"]);row.state="running";row.lease_owner="dead-process";row.lease_until=now()-timedelta(seconds=1)
        db.commit()
    calls=[]
    def fake_review(fields,keywords,username):
        calls.append(fields)
        if len(calls)==1:
            raise TimeoutError("secret-must-not-leak")
        return ReviewVerdict(status="passed",explanation="符合已配置规则"), {"choices":[]}
    monkeypatch.setattr(worker,"review",fake_review)
    assert worker.process_once()
    result=client.get('/api/runs/'+run["id"],headers=h).json()
    assert (result["state"],result["completed"],result["failed"]) == ("partial_failed",3,1)
    raw=client.get('/api/runs/'+run["id"]+'/results',headers=h).text
    assert "secret-must-not-leak" not in raw
    assert not worker.process_once()

def test_only_one_worker_claims_run(client, headers):
    h=headers()
    report=upload(client,h)
    run=client.post('/api/reports/'+report["id"]+'/runs',headers=h).json()
    with ThreadPoolExecutor(max_workers=2) as pool:
        list(pool.map(lambda _:process_once(),range(2)))
    detail=client.get('/api/runs/'+run["id"],headers=h).json()
    assert detail["completed"] == 3
    with Session() as db:
        audits=db.scalars(select(Audit).where(Audit.action=="run.finished",Audit.object_id==run["id"])).all()
        assert len(audits)==1

def test_keyword_snapshot_stays_unchanged(client, headers):
    admin,user=headers("admin"),headers()
    client.post('/api/keyword-requests',headers=admin,json=change())
    report=upload(client,user)
    run=client.post('/api/reports/'+report["id"]+'/runs',headers=user).json()
    client.post('/api/keyword-requests',headers=admin,json=change(action="update",version=1,value="新值"))
    snapshot=client.get('/api/runs/'+run["id"],headers=user).json()["keyword_snapshot"]
    assert snapshot[0]["version"] == 1
    assert snapshot[0]["payload"]["words"] == ["功率"]

def test_result_and_request_access_are_scoped(client, headers):
    admin,user=headers("admin"),headers()
    report=upload(client,admin)
    run=client.post('/api/reports/'+report["id"]+'/runs',headers=admin).json()
    for suffix in ["","/results","/export"]:
        assert client.get('/api/runs/'+run["id"]+suffix,headers=user).status_code == 404
    request=client.post('/api/keyword-requests',headers=admin,json=change()).json()
    assert client.get('/api/keyword-requests',headers=user).json() == []
    assert not any(a["action"]=="keyword.create" for a in client.get('/api/audit',headers=user).json())

def test_live_reviewer_rejects_unapproved_attachment_host(monkeypatch):
    from app.config import settings
    from app.reviewer import review
    monkeypatch.setattr(settings,"app_mode","production")
    monkeypatch.setattr(settings,"attachment_hosts","intranet.example")
    import pytest
    with pytest.raises(ValueError,match="允许列表"):
        review({"附件链接":"http://evil.example/x.jpg"},[],"demo")

def test_worker_rejects_mismatched_rule_version(client, headers, monkeypatch):
    import app.worker as worker
    h=headers()
    report=upload(client,h)
    run=client.post('/api/reports/'+report["id"]+'/runs',headers=h).json()
    monkeypatch.setattr(worker,"rule_version",lambda:"changed-rule")
    def never_review(*args):
        raise AssertionError("mismatched task must never invoke AI")
    monkeypatch.setattr(worker,"review",never_review)
    assert worker.process_once()
    results=client.get('/api/runs/'+run["id"]+'/results',headers=h).json()
    assert all(r["status"]=="error" and r["raw"]["error_type"]=="RuntimeError" for r in results)
