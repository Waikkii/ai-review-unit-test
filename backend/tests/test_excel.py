from io import BytesIO

import pytest
from openpyxl import Workbook

from app.excel import parse_excel
from app.schemas import COLUMNS
from app.worker import process_once


def detail_workbook(header_rows=3, title_rows=2):
    book = Workbook()
    book.active.title = "封面"
    book.active["A1"] = "测试报告"
    book.create_sheet("目录")["A1"] = "领域列表"
    sheet = book.create_sheet("领域1")
    sheet["A1"] = "测试总结"
    sheet["J9"] = "总结中的值不参与审查"
    sheet["A14"] = "测试详情"
    sheet.merge_cells(start_row=14, end_row=13 + title_rows, start_column=1, end_column=3)
    top = 14 + title_rows
    bottom = top + header_rows - 1
    for col, name in enumerate(["测试对象", *COLUMNS[:7]], start=1):
        sheet.cell(top, col, name)
        sheet.merge_cells(start_row=top, end_row=bottom, start_column=col, end_column=col)
    sheet.cell(top, 9, "第一轮")
    sheet.merge_cells(start_row=top, end_row=bottom - 1, start_column=9, end_column=12)
    for col, name in enumerate(COLUMNS[7:], start=9):
        sheet.cell(bottom, col, name)
    sheet.cell(bottom + 1, 4, "供应商名称")
    sheet.merge_cells(start_row=bottom + 1, end_row=bottom + 1, start_column=4, end_column=8)
    sheet.cell(bottom + 1, 9, "/")
    sheet.merge_cells(start_row=bottom + 1, end_row=bottom + 1, start_column=9, end_column=12)
    first = bottom + 2
    sheet.cell(first, 2, "TC-01")
    sheet.cell(first, 1, "设备A")
    for col in (1, 2):
        sheet.merge_cells(start_row=first, end_row=first + 1, start_column=col, end_column=col)
    sheet.cell(first, 10, 0)
    sheet.cell(first + 1, 10, "正常")
    sheet.cell(first + 1, 12, "附件占位")
    sheet.cell(first + 2, 10, "  ")
    sheet.cell(first + 2, 12, "https://intranet.example/skip.jpg")
    sheet.cell(first + 3, 10, "通过")
    sheet.cell(first + 3, 12, "查看附件").hyperlink = "https://intranet.example/ok.jpg"
    return book, sheet


def serialize(book):
    stream = BytesIO()
    book.save(stream)
    book.close()
    return stream.getvalue()


def test_detail_three_row_header_and_value_selection():
    book, _ = detail_workbook()
    rows = parse_excel(serialize(book))
    assert [r["row_number"] for r in rows] == [20, 21, 23]
    assert all(r["domain"] == "领域1" for r in rows)
    assert rows[0]["fields"]["值"] == "0"
    assert rows[0]["fields"]["附件链接"] == ""
    assert rows[1]["fields"]["用例编号"] == "TC-01"
    assert rows[1]["fields"]["测试对象"] == "设备A"
    assert rows[1]["fields"]["附件链接"] == "附件占位"
    assert rows[2]["fields"]["附件链接"] == "https://intranet.example/ok.jpg"


def test_obsolete_two_row_detail_header_is_rejected():
    book, _ = detail_workbook(header_rows=2, title_rows=1)
    with pytest.raises(ValueError, match="三行表头"):
        parse_excel(serialize(book))


def test_three_row_header_rejects_duplicate_value_columns():
    book, sheet = detail_workbook(header_rows=3, title_rows=2)
    sheet["N18"] = "值"
    with pytest.raises(ValueError, match="完整、唯一"):
        parse_excel(serialize(book))


def test_empty_value_skips_other_formulas_but_selected_formula_is_rejected():
    book, sheet = detail_workbook()
    sheet["D22"] = "=1+1"
    assert len(parse_excel(serialize(book))) == 3
    book, sheet = detail_workbook()
    sheet["J20"] = "=1+1"
    with pytest.raises(ValueError, match="公式转成值"):
        parse_excel(serialize(book))


def test_ambiguous_rounds_fail_instead_of_silently_omitting_rows():
    book, sheet = detail_workbook()
    sheet["N18"] = "值"
    with pytest.raises(ValueError, match="完整、唯一"):
        parse_excel(serialize(book))


def test_detail_upload_edit_review_and_export(client, headers):
    book, _ = detail_workbook()
    first = 20
    h = headers()
    response = client.post("/api/reports", headers=h,
                           files={"file": ("detail.xlsx", serialize(book))})
    assert response.status_code == 201, response.text
    report = response.json()
    items = client.get(f'/api/reports/{report["id"]}/items', headers=h).json()
    item = next(i for i in items if i["row_number"] == first)
    edit = {"version": 1, "fields": item["current"], "reason": "修订文字项"}
    assert client.patch(f'/api/items/{item["id"]}', headers=h, json=edit).status_code == 200
    edit["fields"] = {**edit["fields"], "值": " "}
    edit["version"] = 2
    assert client.patch(f'/api/items/{item["id"]}', headers=h, json=edit).status_code == 422
    run = client.post(f'/api/reports/{report["id"]}/runs', headers=h).json()
    assert process_once()
    results = client.get(f'/api/runs/{run["id"]}/results', headers=h).json()
    assert len(results) == 3
    assert all(r["status"] == "needs_review" for r in results)
    from openpyxl import load_workbook
    exported = client.get(f'/api/runs/{run["id"]}/export', headers=h)
    with BytesIO(exported.content) as stream:
        out = load_workbook(stream)
        assert out.sheetnames[:3] == ["封面", "目录", "领域1"]
        assert out["领域1"].cell(first, 10).value == 0
        assert out.worksheets[-1].max_row == 4
        out.close()
